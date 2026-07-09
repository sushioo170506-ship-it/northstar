#!/usr/bin/env python3
"""Wcget Step 6 — 将 report.md 生成 Word 文档。
用法: python3 build_docx.py [gongwen|general]
- gongwen: 参照 GB/T 9704-2012 党政机关公文格式
- general: 通用商务/研报格式
图表引用替换为插入 figures_png/*.png（居中，图题在下）。
"""
import os, re, sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = os.path.join(os.path.dirname(__file__), "..", "output")
MD = os.path.join(BASE, "report.md")
PNG = os.path.join(BASE, "figures_png")

def set_cn_font(run, cn, en, size, bold=False, color=None):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = en
    r = run._element
    rpr = r.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts"); rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), cn)
    rfonts.set(qn("w:ascii"), en)
    rfonts.set(qn("w:hAnsi"), en)
    if color:
        run.font.color.rgb = color

def line_spacing_fixed(p, pts):
    pPr = p._p.get_or_add_pPr()
    spacing = pPr.find(qn("w:spacing"))
    if spacing is None:
        spacing = OxmlElement("w:spacing"); pPr.append(spacing)
    spacing.set(qn("w:line"), str(int(pts * 20)))
    spacing.set(qn("w:lineRule"), "exact")

def strip_inline(text):
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"`(.+?)`", r"\1", text)
    return text

def add_runs_with_bold(p, text, cn, en, size, color=None):
    """按 **bold** 分段添加 run，保留加粗。"""
    parts = re.split(r"(\*\*.+?\*\*)", text)
    for part in parts:
        if not part:
            continue
        b = part.startswith("**") and part.endswith("**")
        content = part[2:-2] if b else part
        content = re.sub(r"`(.+?)`", r"\1", content)
        run = p.add_run(content)
        set_cn_font(run, cn, en, size, bold=b, color=color)

# ---------- 模板配置 ----------
TEMPLATES = {
    "gongwen": dict(
        margins=(3.7, 3.5, 2.8, 2.6),  # 上下左右 cm
        h1_cn="方正小标宋简体", h1_fallback_cn="宋体", h1_en="Times New Roman", h1_size=22,
        body_cn="仿宋_GB2312", body_fallback_cn="FangSong", body_en="Times New Roman", body_size=16,
        sec_cn="黑体", sec_en="Times New Roman", sec_size=16,
        line_pts=29, first_indent=True, sec_color=None, title_color=None,
    ),
    "general": dict(
        margins=(2.54, 2.54, 3.17, 3.17),
        h1_cn="微软雅黑", h1_fallback_cn="微软雅黑", h1_en="Arial", h1_size=21,
        body_cn="宋体", body_fallback_cn="宋体", body_en="Times New Roman", body_size=11,
        sec_cn="微软雅黑", sec_en="Arial", sec_size=15,
        line_pts=None, first_indent=False, sec_color=RGBColor(0x00, 0x5A, 0xB5), title_color=RGBColor(0x0b, 0x11, 0x20),
    ),
}

def build(template):
    cfg = TEMPLATES[template]
    doc = Document()
    sec = doc.sections[0]
    up, dn, lf, rt = cfg["margins"]
    sec.top_margin = Cm(up); sec.bottom_margin = Cm(dn)
    sec.left_margin = Cm(lf); sec.right_margin = Cm(rt)

    # 页码（页脚居中）
    footer = sec.footer
    fp = footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run(); set_cn_font(run, cfg["body_fallback_cn"], cfg["body_en"], 9)
    fld1 = OxmlElement('w:fldSimple'); fld1.set(qn('w:instr'), 'PAGE'); run._element.addprevious(fld1)

    with open(MD, encoding="utf-8") as f:
        lines = f.read().split("\n")

    fig_idx = 0
    i = 0
    while i < len(lines):
        s = lines[i].strip()
        if not s:
            i += 1; continue
        m = re.match(r"!\[(.*?)\]\((.*?)\)", s)
        if m:
            alt, path = m.group(1), m.group(2)
            png = os.path.join(PNG, os.path.basename(path).replace(".svg", ".png"))
            if os.path.exists(png):
                p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(png, width=Inches(5.6))
                cap = doc.add_paragraph(); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = cap.add_run(strip_inline(alt))
                set_cn_font(r, cfg["body_fallback_cn"], cfg["body_en"], cfg["body_size"]-2.5, color=RGBColor(0x66,0x66,0x66))
            i += 1; continue
        if s.startswith("# "):
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_runs_with_bold(p, s[2:], cfg["h1_cn"], cfg["h1_en"], cfg["h1_size"], color=cfg["title_color"])
            for r in p.runs: r.font.bold = True
            if cfg["line_pts"]: line_spacing_fixed(p, cfg["line_pts"]+4)
            doc.add_paragraph()
            i += 1; continue
        if s.startswith("## "):
            p = doc.add_paragraph()
            r = p.add_run(strip_inline(s[3:]))
            set_cn_font(r, cfg["sec_cn"], cfg["sec_en"], cfg["sec_size"], bold=True, color=cfg["sec_color"])
            if cfg["line_pts"]: line_spacing_fixed(p, cfg["line_pts"])
            i += 1; continue
        if s == "---":
            i += 1; continue
        if s.startswith(">"):
            buf = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                buf.append(lines[i].strip()[1:].strip()); i += 1
            p = doc.add_paragraph()
            add_runs_with_bold(p, " ".join(buf), cfg["body_fallback_cn"], cfg["body_en"], cfg["body_size"]-1)
            for r in p.runs: r.font.italic = True
            if cfg["line_pts"]: line_spacing_fixed(p, cfg["line_pts"])
            continue
        if re.match(r"^\d+\.\s", s):
            p = doc.add_paragraph()
            add_runs_with_bold(p, s, cfg["body_fallback_cn"], cfg["body_en"], cfg["body_size"]-1)
            if cfg["line_pts"]: line_spacing_fixed(p, cfg["line_pts"])
            i += 1; continue
        # 普通段落
        p = doc.add_paragraph()
        if cfg["first_indent"]:
            p.paragraph_format.first_line_indent = Pt(cfg["body_size"] * 2)
        add_runs_with_bold(p, s, cfg["body_fallback_cn"], cfg["body_en"], cfg["body_size"])
        if cfg["line_pts"]: line_spacing_fixed(p, cfg["line_pts"])
        i += 1

    out = os.path.join(BASE, f"report.docx" if template == "general" else "report.gongwen.docx")
    doc.save(out)
    print(f"DOCX({template}) ->", out)

if __name__ == "__main__":
    tpl = sys.argv[1] if len(sys.argv) > 1 else "general"
    build(tpl)
