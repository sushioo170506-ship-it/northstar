"""Render the GPT-Live Markdown report to DOCX, PDF, and Feishu-ready Markdown."""

from __future__ import annotations

import html
import base64
import json
import re
import shutil
from pathlib import Path
from urllib.parse import urlparse

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import (
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "outputs" / "gpt-live-application-research"
SOURCE = OUTPUT / "gpt-live应用探索研究报告.md"
DOCX = OUTPUT / "gpt-live应用探索研究报告.docx"
PDF = OUTPUT / "gpt-live应用探索研究报告.pdf"
FEISHU = OUTPUT / "gpt-live应用探索研究报告-飞书导入.md"
MANIFEST = OUTPUT / "gpt-live应用探索研究报告-飞书发布清单.json"
PPTX = OUTPUT / "gpt-live应用探索研究报告-汇报版.pptx"
SLIDES_HTML = OUTPUT / "gpt-live应用探索研究报告-汇报版.html"
PPTX_ASCII = OUTPUT / "gpt-live-deck.pptx"
SLIDES_HTML_ASCII = OUTPUT / "gpt-live-deck.html"
DECK_ZIP = OUTPUT / "gpt-live-deck.zip"

LINK = re.compile(r"\[([^\]]+)\]\((https?://[^)]+)\)")
INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\(https?://[^)]+\))")


def blocks(markdown: str):
    lines = markdown.splitlines()
    index = 0
    paragraph: list[str] = []
    while index < len(lines):
        line = lines[index]
        if paragraph and (
            not line.strip()
            or line.startswith("#")
            or line.startswith("|")
            or line.startswith("```")
            or re.match(r"^\d+\.\s", line)
            or line.startswith("- ")
            or line.startswith("> ")
        ):
            yield "paragraph", " ".join(paragraph)
            paragraph = []
        if not line.strip():
            index += 1
            continue
        if line.startswith("```"):
            language = line[3:].strip()
            content = []
            index += 1
            while index < len(lines) and not lines[index].startswith("```"):
                content.append(lines[index])
                index += 1
            yield "code", (language, "\n".join(content))
            index += 1
            continue
        if line.startswith("|") and index + 1 < len(lines) and re.match(
            r"^\|?\s*:?-{3,}", lines[index + 1]
        ):
            rows = [table_cells(line)]
            index += 2
            while index < len(lines) and lines[index].startswith("|"):
                rows.append(table_cells(lines[index]))
                index += 1
            yield "table", rows
            continue
        heading = re.match(r"^(#{1,6})\s+(.+)$", line)
        if heading:
            yield "heading", (len(heading.group(1)), heading.group(2))
            index += 1
            continue
        if line.startswith("> "):
            yield "quote", line[2:]
            index += 1
            continue
        item = re.match(r"^(\d+\.\s+|-\s+)(.+)$", line)
        if item:
            yield "list", (item.group(1).strip(), item.group(2))
            index += 1
            continue
        paragraph.append(line.strip())
        index += 1
    if paragraph:
        yield "paragraph", " ".join(paragraph)


def table_cells(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def clean(value: str) -> str:
    value = LINK.sub(lambda match: f"{match.group(1)} ({match.group(2)})", value)
    return re.sub(r"(\*\*|`)", "", value)


def mermaid_edges(code: str) -> list[tuple[str, str, str]]:
    labels = {
        match.group(1): match.group(2)
        for line in code.splitlines()
        if (match := re.search(r'(N\d+)\["(.+)"\]', line))
    }
    edges = []
    for line in code.splitlines():
        match = re.search(r"(N\d+)\s+-->(?:\|([^|]+)\|)?\s+(N\d+)", line)
        if match:
            edges.append(
                (
                    labels.get(match.group(1), match.group(1)),
                    match.group(2) or "",
                    labels.get(match.group(3), match.group(3)),
                )
            )
    return edges


def add_docx_runs(paragraph, text: str) -> None:
    cursor = 0
    for match in INLINE.finditer(text):
        paragraph.add_run(text[cursor:match.start()])
        token = match.group(0)
        link = LINK.fullmatch(token)
        if link:
            add_hyperlink(paragraph, link.group(1), link.group(2))
        elif token.startswith("**"):
            paragraph.add_run(token[2:-2]).bold = True
        else:
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
        cursor = match.end()
    paragraph.add_run(text[cursor:])


def add_hyperlink(paragraph, text: str, url: str) -> None:
    relationship = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.extend((color, underline))
    run.append(properties)
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def render_docx(markdown: str) -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Cm(2.2)
    section.left_margin = section.right_margin = Cm(2.4)
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    normal.font.size = Pt(10.5)
    for kind, value in blocks(markdown):
        if kind == "heading":
            level, title = value
            paragraph = document.add_heading(clean(title), level=min(level, 4))
            if level == 1:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif kind in {"paragraph", "quote"}:
            paragraph = document.add_paragraph()
            if kind == "quote":
                paragraph.style = "Quote"
            add_docx_runs(paragraph, value)
        elif kind == "list":
            marker, text = value
            paragraph = document.add_paragraph(
                style="List Number" if marker[0].isdigit() else "List Bullet"
            )
            add_docx_runs(paragraph, text)
        elif kind == "code":
            language, code = value
            if language == "mermaid":
                document.add_heading("结构化流程图", level=3)
                for left, label, right in mermaid_edges(code):
                    table = document.add_table(rows=1, cols=3)
                    table.style = "Light Shading Accent 1"
                    table.cell(0, 0).text = left
                    table.cell(0, 1).text = f"→ {label}" if label else "→"
                    table.cell(0, 2).text = right
                    for cell in (table.cell(0, 0), table.cell(0, 2)):
                        for run in cell.paragraphs[0].runs:
                            run.bold = True
                continue
            paragraph = document.add_paragraph()
            paragraph.style = "No Spacing"
            run = paragraph.add_run((language + "\n" if language else "") + code)
            run.font.name = "Consolas"
            run.font.size = Pt(8)
        elif kind == "table":
            rows = value
            table = document.add_table(rows=len(rows), cols=max(map(len, rows)))
            table.style = "Table Grid"
            for row_index, row in enumerate(rows):
                for column_index, cell in enumerate(row):
                    target = table.cell(row_index, column_index)
                    target.text = clean(cell)
                    if row_index == 0:
                        for run in target.paragraphs[0].runs:
                            run.bold = True
    document.save(DOCX)


def pdf_markup(text: str) -> str:
    text = text.translate(
        str.maketrans(
            {
                "\u2011": "-",
                "\u2013": "-",
                "\u2014": "-",
                "→": "->",
                "⇄": "<->",
                "↓": "v",
            }
        )
    )
    escaped = html.escape(text)
    escaped = re.sub(r"\*\*([^*]+)\*\*", r"<b>\1</b>", escaped)
    escaped = re.sub(r"`([^`]+)`", r"<font name='Courier'>\1</font>", escaped)
    escaped = LINK.sub(
        lambda match: (
            f"<link href='{html.escape(match.group(2))}' color='#0563C1'>"
            f"{html.escape(match.group(1))}</link>"
        ),
        escaped,
    )
    return escaped


def render_pdf(markdown: str) -> None:
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    base = getSampleStyleSheet()
    normal = ParagraphStyle(
        "ChineseBody",
        parent=base["BodyText"],
        fontName="STSong-Light",
        fontSize=9.5,
        leading=15,
        spaceAfter=7,
        wordWrap="CJK",
    )
    quote = ParagraphStyle(
        "ChineseQuote",
        parent=normal,
        leftIndent=12,
        borderColor=colors.HexColor("#D9D9D9"),
        borderWidth=1,
        borderPadding=6,
        backColor=colors.HexColor("#F7F7F7"),
    )
    headings = {
        level: ParagraphStyle(
            f"H{level}",
            parent=normal,
            fontName="STSong-Light",
            fontSize=max(11, 21 - level * 2),
            leading=max(15, 25 - level * 2),
            spaceBefore=12,
            spaceAfter=8,
            alignment=TA_CENTER if level == 1 else 0,
            textColor=colors.HexColor("#17365D"),
        )
        for level in range(1, 7)
    }
    story = []
    for kind, value in blocks(markdown):
        if kind == "heading":
            level, title = value
            if level == 1 and story:
                story.append(PageBreak())
            story.append(Paragraph(pdf_markup(title), headings[level]))
        elif kind == "paragraph":
            story.append(Paragraph(pdf_markup(value), normal))
        elif kind == "quote":
            story.append(Paragraph(pdf_markup(value), quote))
        elif kind == "list":
            marker, text = value
            bullet = marker
            story.append(Paragraph(pdf_markup(text), normal, bulletText=bullet))
        elif kind == "code":
            language, code = value
            if language == "mermaid":
                story.append(Paragraph("结构化流程图", headings[3]))
                for left, label, right in mermaid_edges(code):
                    flow = Table(
                        [
                            [
                                Paragraph(pdf_markup(left), normal),
                                Paragraph(
                                    pdf_markup(f"→ {label}" if label else "→"),
                                    normal,
                                ),
                                Paragraph(pdf_markup(right), normal),
                            ]
                        ],
                        colWidths=[6.2 * cm, 2.0 * cm, 6.2 * cm],
                    )
                    flow.setStyle(
                        TableStyle(
                            [
                                ("BOX", (0, 0), (0, 0), 0.8, colors.HexColor("#2F75B5")),
                                ("BOX", (2, 0), (2, 0), 0.8, colors.HexColor("#2F75B5")),
                                ("BACKGROUND", (0, 0), (0, 0), colors.HexColor("#D9EAF7")),
                                ("BACKGROUND", (2, 0), (2, 0), colors.HexColor("#D9EAF7")),
                                ("ALIGN", (1, 0), (1, 0), "CENTER"),
                                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                                ("TOPPADDING", (0, 0), (-1, -1), 5),
                                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                            ]
                        )
                    )
                    story.extend((flow, Spacer(1, 0.12 * cm)))
                continue
            story.append(
                Paragraph(
                    pdf_markup((language + "\n" if language else "") + code).replace(
                        "\n", "<br/>"
                    ),
                    ParagraphStyle(
                        "Code", parent=normal, fontName="STSong-Light",
                        fontSize=7.5, backColor=colors.HexColor("#F4F4F4"),
                        borderPadding=5,
                    ),
                )
            )
        elif kind == "table":
            rows = [
                [Paragraph(pdf_markup(cell), normal) for cell in row]
                for row in value
            ]
            table = Table(rows, repeatRows=1, hAlign="LEFT")
            table.setStyle(
                TableStyle(
                    [
                        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#B7B7B7")),
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#D9EAF7")),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("LEFTPADDING", (0, 0), (-1, -1), 4),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                        ("TOPPADDING", (0, 0), (-1, -1), 4),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                    ]
                )
            )
            story.extend((table, Spacer(1, 0.2 * cm)))
    document = SimpleDocTemplate(
        str(PDF),
        pagesize=A4,
        rightMargin=1.6 * cm,
        leftMargin=1.6 * cm,
        topMargin=1.7 * cm,
        bottomMargin=1.7 * cm,
        title="GPT-Live应用探索研究报告",
        author="Northstar Research Workflow",
    )
    document.build(story)


def render_feishu(markdown: str) -> None:
    shutil.copyfile(SOURCE, FEISHU)
    tables = sum(kind == "table" for kind, _ in blocks(markdown))
    MANIFEST.write_text(
        json.dumps(
            {
                "format": "feishu-ready-markdown",
                "source": FEISHU.name,
                "native_sheet_candidates": tables,
                "inline_links_preserved": True,
                "publish_status": "pending_user_oauth",
                "note": (
                    "注入FeishuDocumentRenderer后，正文转Docx块，Markdown表格转"
                    "原生Sheet Block；当前环境没有用户飞书凭证，未生成虚假链接。"
                ),
            },
            ensure_ascii=False,
            indent=2,
        )
        + "\n",
        encoding="utf-8",
    )


def render_slides(markdown: str) -> None:
    import sys

    sys.path.insert(0, str(ROOT / "src"))
    from research_workflow.office_renderers import SlidesRenderer

    renderer = SlidesRenderer()
    pptx, _ = renderer.render(
        content=markdown, output_format="pptx", visualizations={"assets": []}
    )
    PPTX.write_bytes(base64.b64decode(pptx))
    PPTX_ASCII.write_bytes(PPTX.read_bytes())
    slides_html, _ = renderer.render(
        content=markdown,
        output_format="slides_html",
        visualizations={"assets": []},
    )
    SLIDES_HTML.write_text(slides_html, encoding="utf-8")
    SLIDES_HTML_ASCII.write_text(slides_html, encoding="utf-8")
    package, _ = renderer.render(
        content=markdown,
        output_format="slides_zip",
        visualizations={"assets": []},
    )
    DECK_ZIP.write_bytes(base64.b64decode(package))


def main() -> None:
    markdown = SOURCE.read_text(encoding="utf-8")
    render_docx(markdown)
    render_pdf(markdown)
    render_feishu(markdown)
    render_slides(markdown)
    print(
        json.dumps(
            {
                "markdown": str(SOURCE),
                "docx": str(DOCX),
                "pdf": str(PDF),
                "feishu": str(FEISHU),
                "manifest": str(MANIFEST),
                "pptx": str(PPTX),
                "slides_html": str(SLIDES_HTML),
                "pptx_ascii": str(PPTX_ASCII),
                "slides_html_ascii": str(SLIDES_HTML_ASCII),
                "deck_zip": str(DECK_ZIP),
            },
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    main()
