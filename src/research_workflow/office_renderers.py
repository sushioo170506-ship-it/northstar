"""Professional DOCX and slide renderers adapted from vetted open-source practices."""

from __future__ import annotations

import base64
import html
import io
import re
from datetime import date
from typing import Any

from .contracts import DocumentRenderer
from .renderers import MarkdownTableParser


def _strip_markup(text: str) -> str:
    text = re.sub(r'<a\s+id="[^"]+"></a>', "", text)
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    return re.sub(r"(\*\*|__|`)", "", text).strip()


def _document_title(content: str) -> str:
    match = re.search(r"(?m)^#\s+(.+)$", content)
    return _strip_markup(match.group(1)) if match else "研究报告"


def _slide_sections(content: str) -> list[tuple[str, str]]:
    title = _document_title(content)
    matches = list(re.finditer(r"(?m)^##\s+(.+)$", content))
    sections = []
    for index, match in enumerate(matches):
        end = matches[index + 1].start() if index + 1 < len(matches) else len(content)
        sections.append((_strip_markup(match.group(1)), content[match.end():end].strip()))
    return [(title, ""), *sections]


class CompositeDocumentRenderer(DocumentRenderer):
    """Route output formats to dedicated renderers."""

    def __init__(self, renderers: dict[str, DocumentRenderer]) -> None:
        self.renderers = dict(renderers)

    def render(self, *, content, output_format, visualizations):
        renderer = self.renderers.get(output_format)
        if renderer is None:
            raise ValueError(f"未配置 {output_format} 对应的文档渲染器")
        return renderer.render(
            content=content,
            output_format=output_format,
            visualizations=visualizations,
        )


class AestheticDocxRenderer(DocumentRenderer):
    """Create a polished, editable DOCX with stable OOXML styles."""

    def render(self, *, content, output_format, visualizations):
        if output_format != "docx":
            raise ValueError(f"AestheticDocxRenderer 不支持: {output_format}")
        try:
            from docx import Document
            from docx.enum.section import WD_SECTION
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            from docx.oxml import OxmlElement
            from docx.oxml.ns import qn
            from docx.shared import Cm, Pt, RGBColor
        except ImportError as exc:
            raise RuntimeError(
                "专业Word渲染需要安装可选依赖: pip install python-docx"
            ) from exc

        document = Document()
        section = document.sections[0]
        section.top_margin = Cm(2.2)
        section.bottom_margin = Cm(2.0)
        section.left_margin = section.right_margin = Cm(2.4)
        section.header_distance = Cm(0.8)
        section.footer_distance = Cm(0.8)

        styles = document.styles
        normal = styles["Normal"]
        normal.font.name = "Arial"
        normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        normal.font.size = Pt(10.5)
        normal.paragraph_format.line_spacing = 1.5
        normal.paragraph_format.space_after = Pt(6)
        colors = ("17365D", "1F4E79", "2F75B5", "5B9BD5")
        for level in range(1, 5):
            style = styles[f"Heading {level}"]
            style.font.name = "Arial"
            style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
            style.font.color.rgb = RGBColor.from_string(colors[level - 1])
            style.font.size = Pt(20 - level * 2)
            style.font.bold = True
            style.paragraph_format.space_before = Pt(12)
            style.paragraph_format.space_after = Pt(6)

        header = section.header.paragraphs[0]
        header.text = "NORTHSTAR RESEARCH  ·  研究报告"
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        header.runs[0].font.size = Pt(8)
        header.runs[0].font.color.rgb = RGBColor(120, 120, 120)
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.add_run("第 ")
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        footer._p.append(field)
        footer.add_run(" 页")

        title = _document_title(content)
        cover = document.add_paragraph()
        cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cover.paragraph_format.space_before = Pt(130)
        run = cover.add_run(title)
        run.bold = True
        run.font.size = Pt(28)
        run.font.color.rgb = RGBColor.from_string("17365D")
        subtitle = document.add_paragraph()
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subtitle.add_run(f"专业研究报告  ·  {date.today().isoformat()}").italic = True
        document.add_page_break()

        toc_title = document.add_paragraph("目录", style="Heading 1")
        toc_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        toc = document.add_paragraph()
        toc_field = OxmlElement("w:fldSimple")
        toc_field.set(qn("w:instr"), 'TOC \\o "1-3" \\h \\z \\u')
        toc._p.append(toc_field)
        document.add_page_break()

        for kind, value in MarkdownTableParser.split_document(content):
            if kind == "markdown":
                in_code = False
                code_lines: list[str] = []
                for line in str(value).splitlines():
                    if line.startswith("# "):
                        continue
                    if line.startswith("```"):
                        if in_code:
                            self._add_code_or_flow(document, code_lines)
                            code_lines = []
                        in_code = not in_code
                        continue
                    if in_code:
                        code_lines.append(line)
                        continue
                    heading = re.match(r"^(#{2,6})\s+(.+)$", line)
                    if heading:
                        document.add_heading(
                            _strip_markup(heading.group(2)),
                            level=min(len(heading.group(1)) - 1, 4),
                        )
                        continue
                    if not line.strip():
                        continue
                    item = re.match(r"^(\s*)(\d+\.\s+|-\s+)(.+)$", line)
                    paragraph = document.add_paragraph(
                        style=(
                            "List Number" if item and item.group(2)[0].isdigit()
                            else "List Bullet" if item else None
                        )
                    )
                    text = item.group(3) if item else line.lstrip("> ")
                    paragraph.add_run(_strip_markup(text))
                continue
            rows = value.values
            table = document.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Light Shading Accent 1"
            table.autofit = True
            for row_index, row in enumerate(rows):
                for column_index, cell in enumerate(row):
                    target = table.cell(row_index, column_index)
                    target.text = cell
                    if row_index == 0:
                        target.paragraphs[0].runs[0].bold = True

        buffer = io.BytesIO()
        document.save(buffer)
        data = buffer.getvalue()
        return (
            base64.b64encode(data).decode("ascii"),
            {
                "rendered": True,
                "renderer": "aesthetic_docx",
                "output_format": "docx",
                "media_type": (
                    "application/vnd.openxmlformats-officedocument."
                    "wordprocessingml.document"
                ),
                "encoding": "base64",
                "byte_count": len(data),
                "editable": True,
                "style_profile": "professional_blue",
                "provenance": [
                    "https://github.com/MiniMax-AI/skills/tree/main/skills/minimax-docx",
                    "https://clawhub.ai/ivangdavila/word-docx",
                ],
            },
        )

    @staticmethod
    def _add_code_or_flow(document, lines: list[str]) -> None:
        if lines and lines[0].strip().startswith("flowchart"):
            labels = {
                match.group(1): match.group(2)
                for line in lines
                if (match := re.search(r'(N\d+)\["(.+)"\]', line))
            }
            document.add_paragraph("结构化流程图", style="Heading 3")
            edge_count = 0
            for line in lines:
                edge = re.search(
                    r"(N\d+)\s+-->(?:\|([^|]+)\|)?\s+(N\d+)", line
                )
                if not edge:
                    continue
                left = labels.get(edge.group(1), edge.group(1))
                right = labels.get(edge.group(3), edge.group(3))
                label = f"〔{edge.group(2)}〕" if edge.group(2) else ""
                paragraph = document.add_paragraph(style="Intense Quote")
                paragraph.add_run(f"{left}  →{label}  {right}")
                edge_count += 1
            if not edge_count:
                document.add_paragraph("流程定义见结构化Mermaid源。", style="Intense Quote")
            return
        paragraph = document.add_paragraph()
        run = paragraph.add_run("\n".join(lines))
        run.font.name = "Consolas"
        run.font.size = 8


class SlidesRenderer(DocumentRenderer):
    """Generate designed fixed-stage HTML slides or editable PPTX."""

    THEME = {
        "bg": "07111f",
        "surface": "10243c",
        "surface2": "163451",
        "text": "F4F8FC",
        "muted": "9AB2C8",
        "cyan": "46D6C8",
        "blue": "55A7FF",
        "amber": "FFBE55",
        "red": "FF6B6B",
    }

    def render(self, *, content, output_format, visualizations):
        if output_format == "slides_html":
            rendered = self._html(content)
            return (
                rendered,
                {
                    "rendered": True,
                    "renderer": "northstar_html_slides",
                    "output_format": "slides_html",
                    "media_type": "text/html",
                    "encoding": "utf-8",
                    "slide_count": len(_slide_sections(content)),
                    "self_contained": True,
                    "provenance": [
                        "https://github.com/zarazhangrui/frontend-slides",
                        "https://github.com/marp-team/marp",
                    ],
                },
            )
        if output_format != "pptx":
            raise ValueError(f"SlidesRenderer 不支持: {output_format}")
        return self._pptx(content)

    @staticmethod
    def _html(content: str) -> str:
        sections = _slide_sections(content)
        slides = [
            SlidesRenderer._html_slide(index, title, body, len(sections))
            for index, (title, body) in enumerate(sections)
        ]
        return """<!doctype html><html lang="zh-CN"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>GPT-Live 应用探索 · 汇报版</title><style>
:root{--bg:#07111f;--surface:#10243c;--surface2:#163451;--text:#f4f8fc;
--muted:#9ab2c8;--cyan:#46d6c8;--blue:#55a7ff;--amber:#ffbe55;--red:#ff6b6b}
*{box-sizing:border-box}html,body{margin:0;width:100%;height:100%;overflow:hidden;
background:#020711;color:var(--text);font-family:"Microsoft YaHei","Noto Sans SC",sans-serif}
#viewport{position:fixed;inset:0;display:flex;align-items:center;justify-content:center}
#stage{width:1920px;height:1080px;position:relative;transform-origin:center center;
background:var(--bg);overflow:hidden;box-shadow:0 30px 100px #000b}
.slide{position:absolute;inset:0;padding:72px 92px 70px;opacity:0;visibility:hidden;
transform:translateY(24px) scale(.992);transition:.45s ease;background:
radial-gradient(circle at 82% 12%,#174b6d88 0,transparent 30%),
linear-gradient(135deg,#07111f 0%,#0a1b2f 68%,#07111f 100%)}
.slide.is-active{opacity:1;visibility:visible;transform:none}.chrome{display:flex;
justify-content:space-between;align-items:center;font:600 18px/1.2 "Arial";letter-spacing:.18em;
color:var(--cyan);text-transform:uppercase}.section-no{color:var(--muted);letter-spacing:.08em}
h1{font-size:58px;line-height:1.12;letter-spacing:-.035em;margin:32px 0 24px;
max-width:1500px}.lead{font-size:27px;line-height:1.55;color:#c8d8e6;max-width:1450px}
.accent{width:112px;height:8px;background:linear-gradient(90deg,var(--cyan),var(--blue));
border-radius:8px;margin:18px 0 30px}.grid{display:grid;gap:24px}.grid-3{grid-template-columns:repeat(3,1fr)}
.grid-2{grid-template-columns:repeat(2,1fr)}.card{background:linear-gradient(145deg,#122a45,#0d2036);
border:1px solid #2a4c69;border-radius:20px;padding:30px;min-height:210px;position:relative}
.card:before{content:"";position:absolute;left:0;top:28px;width:5px;height:48px;
background:var(--cyan);border-radius:0 5px 5px 0}.card h3{font-size:25px;margin:0 0 16px;
color:#fff}.card p,.card li{font-size:21px;line-height:1.52;color:#bed0df}.card ul{margin:0;padding-left:24px}
.metric{font:700 62px/1 "Arial";color:var(--cyan);letter-spacing:-.04em}
.metric-label{font-size:19px;color:var(--muted);margin-top:12px}.metric-card{min-height:170px}
.flow-wrap{display:grid;grid-template-columns:1fr 1fr;gap:8px 34px}.flow{display:grid;
grid-template-columns:1fr 76px 1fr;align-items:center;gap:8px;margin:4px 0}
.node{padding:12px 15px;border:2px solid #3b85bd;background:#102c48;border-radius:10px;
font-size:17px;font-weight:700}.arrow{text-align:center;color:var(--cyan);font-size:19px}
.risk .card:nth-child(2n):before{background:var(--amber)}.risk .card:nth-child(3n):before{background:var(--red)}
.timeline{display:grid;grid-template-columns:160px 1fr;gap:14px 24px;align-items:start}
.step{color:var(--cyan);font:700 24px/1.4 Arial}.step-body{border-left:3px solid #2d5576;
padding:0 0 28px 26px;font-size:22px;line-height:1.5;color:#c4d5e3}
.cover{display:flex;flex-direction:column;justify-content:center}.cover h1{font-size:88px;
max-width:1350px;margin:20px 0}.cover .kicker{color:var(--cyan);font:700 22px Arial;
letter-spacing:.22em}.cover .subtitle{font-size:30px;color:#abc1d4;max-width:1100px}
.cover .orb{position:absolute;width:520px;height:520px;border:2px solid #46d6c855;
border-radius:50%;right:-80px;top:190px;box-shadow:0 0 90px #46d6c829 inset}
.cover .orb:after{content:"LIVE";position:absolute;inset:0;display:grid;place-items:center;
font:800 88px Arial;color:#46d6c822}.page{position:absolute;right:92px;bottom:32px;
font:600 17px Arial;color:var(--muted)}.progress{position:absolute;left:0;bottom:0;height:5px;
background:linear-gradient(90deg,var(--cyan),var(--blue))}.notes{display:none}
#controls{position:fixed;right:24px;bottom:20px;z-index:5;display:flex;gap:8px}
#controls button{width:44px;height:44px;border:1px solid #ffffff44;background:#07111fdd;
color:white;border-radius:50%;font-size:20px;cursor:pointer}.tag{display:inline-block;
padding:8px 14px;border:1px solid #46d6c866;border-radius:999px;color:var(--cyan);
font-size:17px;margin:8px 8px 0 0}
</style></head><body><div id="viewport"><main id="stage">""" + "".join(slides) + """
</main></div><div id="controls"><button onclick="move(-1)">‹</button>
<button onclick="move(1)">›</button></div><script>
const slides=[...document.querySelectorAll('.slide')];let current=0;
function fit(){const s=Math.min(innerWidth/1920,innerHeight/1080);
document.querySelector('#stage').style.transform=`scale(${s})`}
function show(n){slides[current].classList.remove('is-active');current=(n+slides.length)%slides.length;
slides[current].classList.add('is-active');location.hash=`/${current+1}`}
function move(d){show(current+d)}addEventListener('resize',fit);addEventListener('keydown',e=>{
if(['ArrowRight','PageDown',' '].includes(e.key))move(1);
if(['ArrowLeft','PageUp'].includes(e.key))move(-1);if(e.key==='Home')show(0);
if(e.key==='End')show(slides.length-1);if(e.key.toLowerCase()==='f')
document.documentElement.requestFullscreen?.()});fit();
const hash=parseInt(location.hash.replace('#/',''));if(hash>0)show(hash-1);
</script></body></html>"""

    @staticmethod
    def _section_content(body: str) -> dict[str, Any]:
        clean_lines = [
            _strip_markup(line.lstrip("> "))
            for line in body.splitlines()
            if line.strip()
            and not line.startswith(("|", "```", "<a ", "> **表格说明"))
        ]
        subheads = [
            _strip_markup(match.group(1))
            for line in body.splitlines()
            if (match := re.match(r"^###\s+(.+)$", line))
        ]
        bullets = [
            _strip_markup(match.group(1))
            for line in body.splitlines()
            if (match := re.match(r"^(?:\d+\.|[-*])\s+(.+)$", line.strip()))
        ]
        paragraphs = [
            line for line in clean_lines
            if not line.startswith("#") and not re.match(r"^\d+\.\s", line)
        ]
        metrics = []
        for number, suffix in re.findall(
            r"(?<![\w.])(\d+(?:\.\d+)?)(%|亿|万|分钟|秒|ms)?", body
        ):
            display = number + suffix
            if display not in metrics and (
                suffix or float(number) >= 10
            ):
                metrics.append(display)
        tables = [
            value for kind, value in MarkdownTableParser.split_document(body)
            if kind == "table"
        ]
        mermaid = re.search(r"```mermaid\s*\n(.*?)```", body, re.DOTALL)
        edges = []
        if mermaid:
            labels = {
                match.group(1): match.group(2)
                for line in mermaid.group(1).splitlines()
                if (match := re.search(r'(N\d+)\["(.+)"\]', line))
            }
            for line in mermaid.group(1).splitlines():
                edge = re.search(r"(N\d+)\s+-->(?:\|([^|]+)\|)?\s+(N\d+)", line)
                if edge:
                    edges.append(
                        (
                            labels.get(edge.group(1), edge.group(1)),
                            edge.group(2) or "",
                            labels.get(edge.group(3), edge.group(3)),
                        )
                    )
        return {
            "subheads": subheads,
            "bullets": bullets,
            "paragraphs": paragraphs,
            "metrics": metrics,
            "tables": tables,
            "edges": edges,
        }

    @staticmethod
    def _html_slide(index: int, title: str, body: str, total: int) -> str:
        data = SlidesRenderer._section_content(body)
        safe_title = html.escape(title)
        chrome = (
            '<div class="chrome"><span>GPT-LIVE / APPLICATION RESEARCH</span>'
            f'<span class="section-no">{index + 1:02d}</span></div>'
        )
        footer = (
            f'<span class="page">{index + 1} / {total}</span>'
            f'<div class="progress" style="width:{(index + 1) / total * 100:.1f}%"></div>'
        )
        if index == 0:
            return (
                '<section class="slide cover is-active"><div class="orb"></div>'
                '<div class="kicker">STRATEGY · TECHNOLOGY · APPLICATION</div>'
                f"<h1>{safe_title}</h1><div class='accent'></div>"
                "<p class='subtitle'>全双工交互架构、能力边界、竞品格局与场景落地路线</p>"
                "<div><span class='tag'>2026.07</span><span class='tag'>证据驱动</span>"
                "<span class='tag'>决策研究</span></div>"
                "<aside class='notes'>从产品定位切入，重点解释持续交互层与后台推理委托。</aside>"
                f"{footer}</section>"
            )
        header = chrome + f"<h1>{safe_title}</h1><div class='accent'></div>"
        if data["edges"]:
            edges = "".join(
                "<div class='flow'><div class='node'>"
                + html.escape(left)
                + "</div><div class='arrow'>→"
                + (f"<small>{html.escape(label)}</small>" if label else "")
                + "</div><div class='node'>"
                + html.escape(right)
                + "</div></div>"
                for left, label, right in data["edges"][:12]
            )
            content_html = f"<div class='flow-wrap'>{edges}</div>"
        elif len(data["metrics"]) >= 3:
            metric_cards = "".join(
                f"<div class='card metric-card'><div class='metric'>{html.escape(metric)}</div>"
                f"<div class='metric-label'>{html.escape((data['subheads'] + ['关键指标'])[i % max(1, len(data['subheads']) or 1)])}</div></div>"
                for i, metric in enumerate(data["metrics"][:6])
            )
            content_html = f"<div class='grid grid-3'>{metric_cards}</div>"
        else:
            items = data["bullets"] or data["subheads"] or data["paragraphs"]
            items = items[:6]
            cards = "".join(
                f"<article class='card'><h3>{i + 1:02d}</h3>"
                f"<p>{html.escape(item[:150])}</p></article>"
                for i, item in enumerate(items)
            )
            risk = " risk" if "风险" in title else ""
            content_html = f"<div class='grid {'grid-2' if len(items) <= 4 else 'grid-3'}{risk}'>{cards}</div>"
        return (
            f'<section class="slide">{header}{content_html}'
            f"<aside class='notes'>本页围绕“{safe_title}”给出判断和证据。</aside>"
            f"{footer}</section>"
        )

    @staticmethod
    def _pptx(content: str):
        try:
            from pptx import Presentation
            from pptx.dml.color import RGBColor
            from pptx.enum.shapes import MSO_SHAPE
            from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
            from pptx.util import Inches, Pt
        except ImportError as exc:
            raise RuntimeError(
                "PPTX渲染需要安装可选依赖: pip install python-pptx"
            ) from exc
        presentation = Presentation()
        presentation.slide_width = Inches(13.333)
        presentation.slide_height = Inches(7.5)
        sections = _slide_sections(content)
        for index, (title, body) in enumerate(sections):
            slide = presentation.slides.add_slide(
                presentation.slide_layouts[6]
            )
            background = slide.background.fill
            background.solid()
            background.fore_color.rgb = RGBColor.from_string(
                SlidesRenderer.THEME["bg"]
            )
            if index == 0:
                SlidesRenderer._pptx_cover(
                    slide, title, len(sections), MSO_SHAPE, RGBColor, Inches, Pt
                )
                continue
            data = SlidesRenderer._section_content(body)
            accent = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.12), Inches(7.5)
            )
            accent.fill.solid()
            accent.fill.fore_color.rgb = RGBColor.from_string(
                SlidesRenderer.THEME["cyan"]
            )
            accent.line.fill.background()
            eyebrow = slide.shapes.add_textbox(
                Inches(0.72), Inches(0.34), Inches(6.5), Inches(0.3)
            ).text_frame
            eyebrow.text = "GPT-LIVE  /  APPLICATION RESEARCH"
            eyebrow.paragraphs[0].font.size = Pt(9)
            eyebrow.paragraphs[0].font.bold = True
            eyebrow.paragraphs[0].font.color.rgb = RGBColor.from_string(
                SlidesRenderer.THEME["cyan"]
            )
            title_box = slide.shapes.add_textbox(
                Inches(0.72), Inches(0.75), Inches(11.7), Inches(0.72)
            )
            title_frame = title_box.text_frame
            title_frame.text = title
            title_frame.paragraphs[0].font.size = Pt(27)
            title_frame.paragraphs[0].font.bold = True
            title_frame.paragraphs[0].font.color.rgb = RGBColor.from_string(
                SlidesRenderer.THEME["text"]
            )
            title_frame.paragraphs[0].alignment = PP_ALIGN.LEFT
            if data["edges"]:
                SlidesRenderer._pptx_flow(
                    slide, data["edges"][:12], MSO_SHAPE, RGBColor, Inches, Pt,
                    PP_ALIGN, MSO_ANCHOR,
                )
            elif len(data["metrics"]) >= 3:
                SlidesRenderer._pptx_metrics(
                    slide, data, MSO_SHAPE, RGBColor, Inches, Pt
                )
            else:
                SlidesRenderer._pptx_cards(
                    slide, data, "风险" in title, MSO_SHAPE, RGBColor, Inches, Pt
                )
            page = slide.shapes.add_textbox(
                Inches(11.9), Inches(6.92), Inches(0.65), Inches(0.25)
            ).text_frame
            page.text = f"{index + 1:02d} / {len(sections):02d}"
            page.paragraphs[0].font.size = Pt(9)
            page.paragraphs[0].font.color.rgb = RGBColor.from_string(
                SlidesRenderer.THEME["muted"]
            )
        buffer = io.BytesIO()
        presentation.save(buffer)
        data = buffer.getvalue()
        return (
            base64.b64encode(data).decode("ascii"),
            {
                "rendered": True,
                "renderer": "northstar_pptx",
                "output_format": "pptx",
                "media_type": (
                    "application/vnd.openxmlformats-officedocument."
                    "presentationml.presentation"
                ),
                "encoding": "base64",
                "byte_count": len(data),
                "slide_count": len(sections),
                "editable": True,
                "provenance": [
                    "https://github.com/zarazhangrui/frontend-slides",
                    "https://github.com/lewislulu/html-ppt-skill",
                    "https://github.com/marp-team/marp",
                ],
            },
        )

    @staticmethod
    def _pptx_cover(slide, title, total, MSO_SHAPE, RGBColor, Inches, Pt):
        orb = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(9.5), Inches(1.35), Inches(4.4), Inches(4.4)
        )
        orb.fill.background()
        orb.line.color.rgb = RGBColor.from_string(SlidesRenderer.THEME["cyan"])
        orb.line.transparency = 35
        kicker = slide.shapes.add_textbox(
            Inches(0.85), Inches(1.25), Inches(8.2), Inches(0.4)
        ).text_frame
        kicker.text = "STRATEGY · TECHNOLOGY · APPLICATION"
        kicker.paragraphs[0].font.size = Pt(12)
        kicker.paragraphs[0].font.bold = True
        kicker.paragraphs[0].font.color.rgb = RGBColor.from_string(
            SlidesRenderer.THEME["cyan"]
        )
        box = slide.shapes.add_textbox(
            Inches(0.82), Inches(2.0), Inches(9.8), Inches(1.8)
        ).text_frame
        box.text = title
        box.paragraphs[0].font.size = Pt(42)
        box.paragraphs[0].font.bold = True
        box.paragraphs[0].font.color.rgb = RGBColor.from_string(
            SlidesRenderer.THEME["text"]
        )
        subtitle = slide.shapes.add_textbox(
            Inches(0.88), Inches(4.15), Inches(8.4), Inches(0.8)
        ).text_frame
        subtitle.text = "全双工交互架构、能力边界、竞品格局与场景落地路线"
        subtitle.paragraphs[0].font.size = Pt(20)
        subtitle.paragraphs[0].font.color.rgb = RGBColor.from_string(
            SlidesRenderer.THEME["muted"]
        )
        meta = slide.shapes.add_textbox(
            Inches(0.88), Inches(6.45), Inches(5.0), Inches(0.3)
        ).text_frame
        meta.text = f"2026.07  ·  EVIDENCE-DRIVEN  ·  {total} SLIDES"
        meta.paragraphs[0].font.size = Pt(10)
        meta.paragraphs[0].font.color.rgb = RGBColor.from_string(
            SlidesRenderer.THEME["blue"]
        )

    @staticmethod
    def _pptx_cards(slide, data, risk, MSO_SHAPE, RGBColor, Inches, Pt):
        items = data["bullets"] or data["subheads"] or data["paragraphs"]
        items = items[:6] or ["暂无可展示内容"]
        columns = 2 if len(items) <= 4 else 3
        rows = (len(items) + columns - 1) // columns
        card_w = 11.75 / columns
        card_h = min(2.15, 4.95 / rows)
        accents = ["cyan", "amber", "red"] if risk else ["cyan", "blue", "amber"]
        for i, item in enumerate(items):
            col, row = i % columns, i // columns
            x, y = 0.72 + col * (card_w + 0.16), 1.75 + row * (card_h + 0.18)
            shape = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                Inches(x), Inches(y), Inches(card_w), Inches(card_h),
            )
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor.from_string(
                SlidesRenderer.THEME["surface"]
            )
            shape.line.color.rgb = RGBColor.from_string("294B68")
            marker = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                Inches(x), Inches(y + 0.22), Inches(0.06), Inches(0.48),
            )
            marker.fill.solid()
            marker.fill.fore_color.rgb = RGBColor.from_string(
                SlidesRenderer.THEME[accents[i % len(accents)]]
            )
            marker.line.fill.background()
            frame = shape.text_frame
            frame.clear()
            frame.margin_left = Inches(0.25)
            frame.margin_right = Inches(0.2)
            frame.margin_top = Inches(0.18)
            p = frame.paragraphs[0]
            p.text = f"{i + 1:02d}"
            p.font.size = Pt(12)
            p.font.bold = True
            p.font.color.rgb = RGBColor.from_string(
                SlidesRenderer.THEME[accents[i % len(accents)]]
            )
            body = frame.add_paragraph()
            body.text = item[:120]
            body.font.size = Pt(15 if len(item) < 80 else 12)
            body.font.color.rgb = RGBColor.from_string(
                SlidesRenderer.THEME["text"]
            )
            body.space_before = Pt(8)

    @staticmethod
    def _pptx_metrics(slide, data, MSO_SHAPE, RGBColor, Inches, Pt):
        metrics = data["metrics"][:6]
        for i, metric in enumerate(metrics):
            col, row = i % 3, i // 3
            x, y = 0.72 + col * 4.08, 1.8 + row * 2.35
            shape = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                Inches(x), Inches(y), Inches(3.75), Inches(1.95),
            )
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor.from_string(
                SlidesRenderer.THEME["surface"]
            )
            shape.line.color.rgb = RGBColor.from_string("2C5675")
            frame = shape.text_frame
            frame.clear()
            p = frame.paragraphs[0]
            p.text = metric
            p.font.size = Pt(30)
            p.font.bold = True
            p.font.color.rgb = RGBColor.from_string(
                SlidesRenderer.THEME["cyan"]
            )
            label = frame.add_paragraph()
            labels = data["subheads"] or ["关键证据"]
            label.text = labels[i % len(labels)][:45]
            label.font.size = Pt(12)
            label.font.color.rgb = RGBColor.from_string(
                SlidesRenderer.THEME["muted"]
            )
            label.space_before = Pt(8)

    @staticmethod
    def _pptx_flow(
        slide, edges, MSO_SHAPE, RGBColor, Inches, Pt, PP_ALIGN, MSO_ANCHOR
    ):
        for i, (left, label, right) in enumerate(edges):
            row, col = i % 6, i // 6
            base_x = 0.72 + col * 6.15
            y = 1.62 + row * 0.84
            for x, text in ((base_x, left), (base_x + 3.5, right)):
                shape = slide.shapes.add_shape(
                    MSO_SHAPE.ROUNDED_RECTANGLE,
                    Inches(x), Inches(y), Inches(2.62), Inches(0.62),
                )
                shape.fill.solid()
                shape.fill.fore_color.rgb = RGBColor.from_string(
                    SlidesRenderer.THEME["surface2"]
                )
                shape.line.color.rgb = RGBColor.from_string(
                    SlidesRenderer.THEME["blue"]
                )
                frame = shape.text_frame
                frame.text = text[:36]
                frame.vertical_anchor = MSO_ANCHOR.MIDDLE
                p = frame.paragraphs[0]
                p.alignment = PP_ALIGN.CENTER
                p.font.size = Pt(11)
                p.font.bold = True
                p.font.color.rgb = RGBColor.from_string(
                    SlidesRenderer.THEME["text"]
                )
            arrow = slide.shapes.add_textbox(
                Inches(base_x + 2.68), Inches(y + 0.10), Inches(0.76), Inches(0.34)
            ).text_frame
            arrow.text = "→" + (f" {label}" if label else "")
            arrow.paragraphs[0].alignment = PP_ALIGN.CENTER
            arrow.paragraphs[0].font.size = Pt(11)
            arrow.paragraphs[0].font.color.rgb = RGBColor.from_string(
                SlidesRenderer.THEME["cyan"]
            )
